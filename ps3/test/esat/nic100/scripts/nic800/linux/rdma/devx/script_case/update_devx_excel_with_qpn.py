#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""在现有合并 Excel 中追加 Reserved QPN 14 条用例。

本脚本读取 ps3/test/esat/nic100/scripts/nic800/linux/rdma/devx/script_case/
目录下已生成的 14 个 QPN 测试脚本，提取 case_id、case_title、test_steps、
test_precondition 等字段，追加到合并 Excel 的现有 PD/QP/CQ/SRQ 用例之后。
"""
import argparse
import ast
import logging
from pathlib import Path
import re
import shutil
from datetime import datetime

import openpyxl


logging.basicConfig(level=logging.INFO, format='%(levelname)s: %(message)s')
logger = logging.getLogger(__name__)

QPN_FILE_GLOB = 'nic_linux_rdma_devx_f_06_53_005_*.py'
EXPECTED_SCRIPT_COUNT = 14

DEFAULT_PRECONDITION = (
    '1. RDMA 设备（mlx5_0）在线且 pyverbs mlx5 可用；\n'
    '2. 已部署 ipvbs 运行时（ipvbs_s.py / ipvbs.py）；\n'
    '3. 测试机与 server 节点 SSH 免密连通。'
)

DEFAULT_EXPECTED_RESULT = (
    '1. CREATE_GENERAL_OBJECT + MLX5_OBJ_TYPE_RESERVED_QPN 命令下发成功；\n'
    '2. 返回有效的 first_qpn，多个 block 的 first_qpn 不重复；\n'
    '3. Reserved QPN block 可正常释放，资源无泄漏。'
)

BASE_FIELDS = {
    '产品型号': '800G网卡',
    '模块': 'rdma',
    '优先级': 'A2',
    '功能等级': 'f',
    '测试类型': '系统测试-功能',
    '测试方法': '系统测试-功能',
    '适用平台': 'Linux',
    '命令': 'devx_obj_create / SXE3R_DEVX_OP_CREATE_GENERAL_OBJECT + MLX5_OBJ_TYPE_RESERVED_QPN',
    '备注': '主机PF; RDMA用户态; Reserved QPN',
}


REPO_ROOT_MARKERS = ('AGENTS.md', '.git')


def _find_repo_root(start: Path | None = None) -> Path:
    """向上查找包含 AGENTS.md 或 .git 的目录作为项目根目录。"""
    path = (start or Path(__file__).resolve()).parent
    while True:
        if any((path / marker).exists() for marker in REPO_ROOT_MARKERS):
            return path
        parent = path.parent
        if parent == path:
            raise FileNotFoundError(
                '无法定位项目根目录：未找到 AGENTS.md 或 .git 标记。'
            )
        path = parent


REPO_ROOT = _find_repo_root()
EXCEL_PATH = REPO_ROOT / 'excel测试用例文件夹' / 'DevX_PD_QP_CQ_SRQ_创建测试用例_含指令.xlsx'
SCRIPT_DIR = Path(__file__).resolve().parent


def _string_value(node: ast.AST | None) -> str | None:
    """从 AST 节点中提取可确定的常量字符串值。"""
    if node is None:
        return None
    if isinstance(node, ast.Constant) and isinstance(node.value, str):
        return node.value
    if isinstance(node, ast.Tuple):
        parts = []
        for elt in node.elts:
            val = _string_value(elt)
            if val is None:
                return None
            parts.append(val)
        return ''.join(parts)
    return None


def _find_target_class(tree: ast.AST, path: Path) -> ast.ClassDef | None:
    """定位脚本中的目标类（NIC100 测试类）。

    优先返回类名以 ``NicLinuxRdmaDevx`` 开头的类；
    若未找到，则返回第一个继承 ``ScriptBase`` 的类。
    """
    classes = [node for node in ast.walk(tree) if isinstance(node, ast.ClassDef)]
    for node in classes:
        if node.name.startswith('NicLinuxRdmaDevx'):
            return node
    for node in classes:
        for base in node.bases:
            base_name = None
            if isinstance(base, ast.Name):
                base_name = base.id
            elif isinstance(base, ast.Attribute):
                base_name = base.attr
            if base_name == 'ScriptBase':
                return node
    logger.warning(f'{path.name}: 未找到继承 ScriptBase 的脚本类，类属性提取可能不准确')
    return None


def _extract_class_string(source: str, attr_name: str, path: Path) -> str | None:
    """使用 AST 从目标类属性赋值中提取常量字符串值。

    支持：
      - 普通字符串字面量
      - 括号包裹的字符串元组/隐式拼接
    对 f-字符串、变量、非字符串表达式返回 None。
    """
    try:
        tree = ast.parse(source)
    except SyntaxError as exc:
        logger.warning(f'解析源码失败: {exc}')
        return None

    target_class = _find_target_class(tree, path)
    if target_class is None:
        return None

    for item in target_class.body:
        targets = []
        value = None
        if isinstance(item, ast.Assign):
            targets = item.targets
            value = item.value
        elif isinstance(item, ast.AnnAssign):
            targets = [item.target]
            value = item.value
        for target in targets:
            if isinstance(target, ast.Name) and target.id == attr_name:
                return _string_value(value)
    return None


def _extract_description(source: str) -> str | None:
    """从模块 docstring 的 description: 行提取描述文本。"""
    match = re.search(
        r'description:\s*(.+?)(?:\n\n|\n\s+@|\n\s*\*\s*@)',
        source,
        re.DOTALL,
    )
    if not match:
        return None
    return ' '.join(line.strip() for line in match.group(1).strip().splitlines())


def _extract_scene(path: Path, source: str) -> str:
    """提取用例场景关键字。

    优先从文件名 ``devx_create_qpn_<scene>_create_qpn.py`` 解析；
    无法解析时回退到 docstring 的 ``description:`` 块第一行；
    两者均失败时返回空字符串并记录警告。
    """
    filename_match = re.search(
        r'devx_create_qpn_(.+?)_create_qpn\.py$',
        path.name,
        re.IGNORECASE,
    )
    if filename_match:
        return filename_match.group(1).replace('_', ' ')

    desc_match = re.search(
        r'description:\s*(.+?)(?:\n\n|\n\s+opcode:|\n\s*\*\s*opcode:)',
        source,
        re.DOTALL,
    )
    if desc_match:
        return desc_match.group(1).strip().split('\n')[0].strip()

    logger.warning(
        f'{path.name}: 无法从文件名或 description 注释中提取场景关键字，'
        '预期结果将使用默认模板'
    )
    return ''


def _case_id_from_filename(path: Path) -> str | None:
    """从脚本文件名推断预期的 case_id。"""
    match = re.search(
        r'^(nic_linux_rdma_devx_f_06_53_005_\d{3})',
        path.stem,
    )
    return match.group(1) if match else None


def _extract_script_fields(path: Path) -> dict:
    """从单个 QPN 脚本中提取用例字段。"""
    source = path.read_text(encoding='utf-8')

    case_id = _extract_class_string(source, 'case_id', path)
    case_title = _extract_class_string(source, 'case_title', path)
    test_steps = _extract_class_string(source, 'test_steps', path)
    test_precondition = _extract_class_string(source, 'test_precondition', path)
    test_object = _extract_class_string(source, 'test_object', path)
    test_module = _extract_class_string(source, 'test_module', path)
    test_tool = _extract_class_string(source, 'test_tool', path)
    case_type = _extract_class_string(source, 'case_type', path)

    # 若脚本未定义 test_precondition，从 description 注释中尝试补充
    if not test_precondition:
        desc_precondition = _extract_description(source)
        if desc_precondition:
            test_precondition = desc_precondition
        else:
            test_precondition = DEFAULT_PRECONDITION

    scene = _extract_scene(path, source)

    # 文件名中的 case_id 与源码中的 case_id 一致性校验
    filename_case_id = _case_id_from_filename(path)
    if filename_case_id and case_id and filename_case_id != case_id:
        logger.warning(
            f'{path.name}: 文件名推断的 case_id ({filename_case_id}) '
            f'与源码 case_id ({case_id}) 不一致'
        )

    return {
        'case_id': case_id or path.stem,
        'case_title': case_title or '',
        'test_steps': test_steps or '',
        'test_precondition': test_precondition,
        'test_object': test_object or 'RDMA Reserved QPN（预留队列对编号）',
        'test_module': test_module or 'rdma',
        'test_tool': test_tool or 'pyverbs',
        'case_type': case_type or '系统测试-功能',
        'scene': scene,
    }


def _sorted_qpn_scripts() -> list[Path]:
    """返回按用例编号排序的 QPN 脚本路径列表。"""
    scripts = sorted(SCRIPT_DIR.glob(QPN_FILE_GLOB))
    # 过滤掉本脚本自身（匹配 glob 不会命中，但保险起见）
    scripts = [p for p in scripts if p.name != 'update_devx_excel_with_qpn.py']
    return scripts


def _build_expected_result(fields: dict) -> str:
    """根据用例场景生成预期结果描述。"""
    scene = fields.get('scene', '').lower()
    if any(k in scene for k in ('异常', 'abnormal', 'anomaly')):
        return (
            '1. 错误 obj_type、超范围 log_obj_range 或非 CREATE_GENERAL_OBJECT opcode 的请求\n'
            '   均被固件拒绝，不返回有效 first_qpn；\n'
            '2. 资源未发生泄漏；\n'
            '3. 环境可正常清理。'
        )
    if any(
        k in scene
        for k in ('压力', 'stress', '稳定性', 'stability', '交替', 'alternating')
    ):
        return (
            '1. 大量/反复创建 Reserved QPN block 均成功；\n'
            '2. 所有 first_qpn 唯一且无重复；\n'
            '3. 释放后资源可再次分配，无泄漏。'
        )
    return DEFAULT_EXPECTED_RESULT


def _build_rows() -> list[tuple[int, dict, Path]]:
    """构建待追加的行数据，返回 (序号, 字段, 脚本路径) 列表。"""
    scripts = _sorted_qpn_scripts()
    if len(scripts) != EXPECTED_SCRIPT_COUNT:
        raise RuntimeError(
            f'期望找到 {EXPECTED_SCRIPT_COUNT} 个 QPN 脚本（glob={QPN_FILE_GLOB}），'
            f'实际找到 {len(scripts)} 个。请先生成缺失的脚本后再运行本工具。'
        )

    rows = []
    for idx, script_path in enumerate(scripts, start=1):
        fields = _extract_script_fields(script_path)
        rows.append((idx, fields, script_path))
    return rows


def _print_row_preview(rows: list[tuple[int, dict, Path]], start_row: int) -> None:
    """打印待追加行的预览。"""
    for idx, fields, script_path in rows:
        row = start_row + idx
        logger.info(f'即将追加第 {row} 行: {fields["case_id"]} ({script_path.name})')
        logger.info(f'  标题: {fields["case_title"][:60]}...')
        logger.info(f'  前置条件: {fields["test_precondition"][:60]}...')
        logger.info(f'  测试步骤: {fields["test_steps"][:60]}...')
        logger.info(f'  预期结果: {_build_expected_result(fields)[:60]}...')


def _backup_excel() -> Path:
    """创建原始 Excel 的时间戳备份。"""
    timestamp = datetime.now().strftime('%Y%m%d_%H%M%S')
    backup_path = EXCEL_PATH.parent / f'{EXCEL_PATH.name}.bak.{timestamp}'
    shutil.copy2(EXCEL_PATH, backup_path)
    logger.info(f'已创建 Excel 备份: {backup_path}')
    return backup_path


def append_qpn_cases(dry_run: bool = False, backup: bool = True) -> int:
    """追加 QPN 用例到 Excel，返回追加行数。

    Args:
        dry_run: 为 True 时只打印预览，不修改文件。
        backup: 为 True 时在写入前创建时间戳备份。
    """
    wb = openpyxl.load_workbook(EXCEL_PATH)
    ws = wb.active
    max_row = ws.max_row

    rows = _build_rows()

    if dry_run:
        logger.info(
            f'--dry-run: 将追加 {len(rows)} 条 QPN 用例到 {EXCEL_PATH}'
            f'（起始行 {max_row + 1}）'
        )
        _print_row_preview(rows, max_row)
        return 0

    if backup:
        _backup_excel()

    appended = 0
    for idx, fields, script_path in rows:
        row = max_row + idx

        # 按现有 Excel 列顺序填充
        ws.cell(row=row, column=1, value=BASE_FIELDS['产品型号'])
        ws.cell(row=row, column=2, value=fields['test_module'])
        ws.cell(row=row, column=3, value=BASE_FIELDS['优先级'])
        ws.cell(row=row, column=4, value=BASE_FIELDS['功能等级'])
        ws.cell(row=row, column=5, value=script_path.name)
        ws.cell(row=row, column=6, value=fields['case_id'])
        ws.cell(row=row, column=7, value=fields['case_title'])
        ws.cell(row=row, column=8, value=fields['test_precondition'])
        ws.cell(row=row, column=9, value=fields['test_steps'])
        ws.cell(row=row, column=10, value=_build_expected_result(fields))
        ws.cell(row=row, column=11, value=fields['case_title'])
        ws.cell(row=row, column=12, value=fields['case_type'])
        ws.cell(row=row, column=13, value=BASE_FIELDS['测试方法'])
        ws.cell(row=row, column=14, value=None)
        ws.cell(row=row, column=15, value=None)
        ws.cell(row=row, column=16, value=BASE_FIELDS['适用平台'])
        ws.cell(row=row, column=17, value=BASE_FIELDS['命令'])
        ws.cell(row=row, column=18, value=BASE_FIELDS['备注'])
        appended += 1

    wb.save(EXCEL_PATH)
    logger.info(f'已在 {EXCEL_PATH} 追加 {appended} 条 QPN 用例')
    return appended


def _parse_args() -> argparse.Namespace:
    parser = argparse.ArgumentParser(
        description='在现有合并 Excel 中追加 Reserved QPN 用例。'
    )
    parser.add_argument(
        '--dry-run',
        action='store_true',
        help='只打印待追加行的预览，不修改 Excel 文件。',
    )
    parser.add_argument(
        '--backup',
        action=argparse.BooleanOptionalAction,
        default=True,
        help='写入前是否创建时间戳备份（默认开启）。',
    )
    return parser.parse_args()


if __name__ == '__main__':
    args = _parse_args()
    append_qpn_cases(dry_run=args.dry_run, backup=args.backup)
