# -*- encoding=utf8 -*-
"""
author: luyafei
date  : 2020.04.21
descriptor: 用Openpyxl实现excel文件的相关功能实现
"""
import os
import sys
import subprocess
import re
import openpyxl
from openpyxl import Workbook
from openpyxl.styles import PatternFill
from openpyxl.worksheet.worksheet import Worksheet
from openpyxl.cell.cell import ILLEGAL_CHARACTERS_RE

from libs import path_libs
from scripts.script_libs import log
from libs import constants


def get_work_book(file_name: str) -> Workbook:
    """
    author: DuPanPan
    date: 2020.04.21
    description: 获取工作簿
    :return work_book:excel文件对应工作簿
    """
    work_book = None
    try:
        work_book = openpyxl.load_workbook(file_name)
    except Exception as result:
        log.error("Not found %s" % result)
        exit(1)
    return work_book


def create_excel_file(columns: list, file_path_name: str, sheet_name: str) -> None:
    """
    author: luyafei
    date  : 2020.04.21
    description: 创建excel文件
    :param columns: excel文件的列
    :param file_path_name: excel文件的路径及文件名
    :param sheet_name: sheet名称
    """
    try:
        workbook = Workbook()
        # 创建带有index的sheet
        sheet = workbook.create_sheet(title=sheet_name, index=0)
        # 将每列的title加入sheet中
        sheet.append(columns)
        # 设置excel表的列宽
        sheet.sheet_format.defaultColWidth = constants.EXCEL_COL_WIDTH
        # 保存excel文件
        workbook.save(file_path_name)
        log.info("make %s success" % file_path_name)
    except Exception as result:
        log.error("create excel file failed. Error is %s" % result)
        exit(1)


def append_info_to_excel(workbook: Workbook, sheet: Worksheet, row_info: list,
                         excel_file: str) -> None:
    """
    author: luyafei
    date  : 2020.04.21
    description: 将每行的数据信息追加到excel文件中
    :param workbook: excel文件的workbook
    :param sheet: 当前操作的sheet
    :param row_info: 要追加的行属性信息
    :param excel_file: 被追加的excel文件
    """
    try:
        sheet.append(row_info)
    except Exception as e:
        log.info(row_info)
        for i in range(len(row_info)):
            row_info[i] = ILLEGAL_CHARACTERS_RE.sub(r'', str(row_info[i]))
        log.info("--------------------")
        log.info(row_info)
        sheet.append(row_info)

    workbook.save(excel_file)


def __add_color(wb_sheet: Worksheet, wb: Workbook, excel_file: str) -> None:
    """
    author: luyafei
    date  : 2020.04.21
    description: 给sheet页中的某一行着色
    :param wb_sheet: sheet页
    :param wb: excel文件的workbook
    """
    fill_green = PatternFill(constants.CELL_FILL_TYPE,
                             fgColor=constants.CELL_COLOR)
    row_num = wb_sheet.max_row
    for col in range(1, wb_sheet.max_column + 1):
        cell_obj = wb_sheet.cell(row=row_num, column=col)
        cell_obj.fill = fill_green
    wb.save(excel_file)


def merge_excels(report_excel_file: str, excel_list: list) -> None:
    """
    author: luyafei
    date  : 2020.04.21
    description: 合并多个excel文件到一个excel文件中
    :param report_excel_file: 测试报告excel文件
    :param excel_list: excel文件列表
    """
    report_wb = openpyxl.load_workbook(report_excel_file)
    report_wb_sheet = report_wb.get_sheet_by_name(constants.REPORT_SHEET)
    # 记录测试报告的文件
    report_bug_file = None

    for excel_file in excel_list:
        title_list = []
        if excel_file.find(constants.BUG) >= 0:
            report_bug_file = excel_file
            title_list.append(constants.BUG_SHEET)
        else:
            title_list.append(constants.RESULT_SHEET)
        append_info_to_excel(report_wb, report_wb_sheet, title_list,
                             report_excel_file)
        # 给添加的bug list和result list行着色
        __add_color(report_wb_sheet, report_wb, report_excel_file)
        excel_file_wb = openpyxl.load_workbook(excel_file)
        excel_file_sheet_names = excel_file_wb.get_sheet_names()
        excel_file_sheet_name = excel_file_wb[excel_file_sheet_names[0]]
        for row in range(1, excel_file_sheet_name.max_row + 1):
            cell_list = []
            for col in range(1, excel_file_sheet_name.max_column + 1):
                cell_list.append(excel_file_sheet_name.cell(row, col).value)
            append_info_to_excel(report_wb, report_wb_sheet, cell_list,
                                 report_excel_file)
    # 删除测试报告的问题列表文件
    if report_bug_file is not None:
        os.remove(report_bug_file)


def get_excel_data_by_row(excel_file: str) -> list:
    """
    author: luyafei
    date  : 20200511
    description: 按行获取excel文件中的数据信息
    :param excel_file: excel文件
    :return: excel文件中第一个sheet中的所有数据信息
    """
    excel_wb = openpyxl.load_workbook(excel_file)
    # 获取表格的所有sheet以及当前要操作的sheet
    sheets = excel_wb.get_sheet_names()
    first_sheet = sheets[0]
    work_sheet = excel_wb[first_sheet]
    max_row_num = work_sheet.max_row
    max_col_num = work_sheet.max_column
    row_data_list = []
    # 获取表格中第一行是title信息，从第二行开始是数据信息
    for row_num in range(constants.CASE_SCRIPT_READ_MIN_ROW, max_row_num + 1):
        row_data_dict = {}
        for col in range(1, max_col_num + 1):
            row_data_dict.setdefault(work_sheet.cell(1, col).value,
                                     work_sheet.cell(row_num, col).value)
        row_data_list.append(row_data_dict)
    return row_data_list


def generate_report_bug_list(bug_column_list: list,
                             report_bug_column_list: list,
                             bug_excel_file: str) -> list:
    """
    author: luyafei
    date  : 20200511
    description: 生成测试报告需要的问题列表
    :param bug_column_list: 测试的问题列表信息
    :param report_bug_column_list: 测试报告的问题列信息
    :param bug_excel_file: 测试生成的所有bug列表的excel文件名称及路径信息
    :return: 测试报告需要的问题列表
    """
    # 存储测试报告问题列信息在测试生成的所有bug列表信息中的索引信息
    report_bug_column_index_list = []
    # 获取测试报告问题列信息在测试生成的所有bug列表信息中的索引信息
    for column_name in report_bug_column_list:
        col_index = bug_column_list.index(column_name)
        report_bug_column_index_list.append(col_index)

    excel_wb = openpyxl.load_workbook(bug_excel_file)
    # 获取表格的所有sheet以及当前要操作的sheet
    sheets = excel_wb.get_sheet_names()
    first_sheet = sheets[0]
    work_sheet = excel_wb[first_sheet]
    # 存储测试报告bug信息列表
    report_bug_list = []
    # 第一行是列标题
    for row in range(constants.CASE_SCRIPT_READ_MIN_ROW,
                     work_sheet.max_row + 1):
        bug_info_list = []
        for col in report_bug_column_index_list:
            # index方法获取到的位置信息从0开始的，而获取excel表格的单元格值时
            # 需要从1开始的
            bug_info_list.append(work_sheet.cell(row, col + 1).value)

        report_bug_list.append(bug_info_list)
    return report_bug_list
