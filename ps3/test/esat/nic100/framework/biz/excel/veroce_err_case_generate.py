#!/usr/bin/env python3
# -*- encoding=utf8 -*-
"""
author: agent
date  : 2026.06.18
description: 根据《veroce错误处理.xlsx》生成符合《多PF测试-测试用例.xlsx》格式的测试用例 Excel。
"""
import os
import re
import openpyxl
from openpyxl import Workbook
from openpyxl.cell.cell import ILLEGAL_CHARACTERS_RE

# 项目根目录（脚本位于 ps3/test/esat/nic100/framework/biz/excel/）
SCRIPT_DIR = os.path.dirname(os.path.abspath(__file__))
PROJECT_ROOT = os.path.abspath(
    os.path.join(SCRIPT_DIR, "..", "..", "..", "..", "..", "..", "..")
)

SOURCE_FILE = os.path.join(PROJECT_ROOT, "veroce错误处理.xlsx")
SOURCE_SHEET = "错误处理"
TARGET_FILE = os.path.join(PROJECT_ROOT, "veroce错误处理-测试用例.xlsx")
TARGET_SHEET = "测试用例"

# 与《多PF测试-测试用例.xlsx》完全一致的列头
TARGET_COLUMNS = [
    "测试编号",
    "用例标题",
    "所属产品",
    "所属模块",
    "优先级",
    "用例类型",
    "脚本名称",
    "前提条件",
    "测试平台实现自动",
    "作者",
    "测试步骤",
    "预期结果",
    "检查项 / 测试点",
    "卡类别",
    "物理环境",
    "测试场景",
    "冒烟测试",
    "最小用例集",
    "分类",
    "关键词",
    "描述",
    "已调试",
    "测试分类集",
    "来源 Bug",
    "用例测试分类",
    "产品发布优先级",
]


def clean_value(value):
    """将单元格值转为字符串并过滤非法 Excel 字符。"""
    if value is None:
        return ""
    text = str(value)
    return ILLEGAL_CHARACTERS_RE.sub(r"", text)


def extract_precondition(steps_text: str) -> str:
    """
    从测试场景/步骤文本中提取基线准备部分作为前提条件。
    启发式规则：取前两条非空步骤（以换行分隔），并去掉常见序号前缀。
    """
    lines = [line.strip() for line in steps_text.splitlines() if line.strip()]
    precondition_lines = []
    for line in lines[:2]:
        # 去掉常见序号/步骤前缀
        line = re.sub(r"^步骤\s*\d+[:：]", "", line)
        line = re.sub(r"^\d+[:：、.]\s*", "", line).strip()
        if line:
            precondition_lines.append(line)
    return "\n".join(precondition_lines)


def build_test_case(index: int, row_data: dict) -> list:
    """将源表一行数据映射为目标测试用例的一行。"""
    err_name = clean_value(row_data.get("错误名称"))
    description = clean_value(row_data.get("描述"))
    scenario = clean_value(row_data.get("测试场景/测试步骤"))
    phenomena = clean_value(row_data.get("系统层面上可以观察到的现象"))

    case_id = f"TC-{index:03d}"
    title = f"{err_name} 错误处理验证" if err_name else f"{case_id} 错误处理验证"
    precondition = extract_precondition(scenario)
    check_point = f"{err_name} - {description}" if err_name or description else ""

    return [
        case_id,
        title,
        "",
        "Veroce错误处理",
        "A1",
        "系统测试-异常",
        "",
        precondition,
        "SNPY",
        "",
        scenario,
        phenomena,
        check_point,
        "",
        "",
        err_name,
        "",
        "",
        "",
        err_name,
        description,
        "",
        "",
        "",
        "",
        "",
    ]


def read_source_rows() -> list:
    """读取源 Excel 的“错误处理”sheet，返回每行的字典列表。"""
    if not os.path.exists(SOURCE_FILE):
        raise FileNotFoundError(f"Source file not found: {SOURCE_FILE}")

    wb = openpyxl.load_workbook(SOURCE_FILE, data_only=True)
    if SOURCE_SHEET not in wb.sheetnames:
        raise ValueError(f"Sheet '{SOURCE_SHEET}' not found in {SOURCE_FILE}")

    ws = wb[SOURCE_SHEET]
    headers = [clean_value(ws.cell(1, col).value) for col in range(1, ws.max_column + 1)]

    rows = []
    for row_idx in range(2, ws.max_row + 1):
        row_data = {}
        has_data = False
        for col_idx, header in enumerate(headers, start=1):
            value = ws.cell(row_idx, col_idx).value
            row_data[header] = value
            if value is not None:
                has_data = True
        if has_data:
            rows.append(row_data)
    return rows


def create_target_excel(test_cases: list) -> None:
    """创建目标测试用例 Excel 文件。"""
    wb = Workbook()
    # 删除 Workbook 默认创建的空白 sheet
    if "Sheet" in wb.sheetnames:
        wb.remove(wb["Sheet"])
    ws = wb.create_sheet(title=TARGET_SHEET, index=0)
    ws.append(TARGET_COLUMNS)
    for case in test_cases:
        ws.append(case)
    # 默认列宽与框架风格保持一致
    ws.sheet_format.defaultColWidth = 20
    wb.save(TARGET_FILE)


def main() -> None:
    source_rows = read_source_rows()
    test_cases = []
    for idx, row in enumerate(source_rows, start=1):
        # 跳过没有错误名称或场景的空行
        if not row.get("错误名称") and not row.get("测试场景/测试步骤"):
            continue
        test_cases.append(build_test_case(idx, row))

    create_target_excel(test_cases)
    print(f"Generated {TARGET_FILE} with {len(test_cases)} test cases.")


if __name__ == "__main__":
    main()
